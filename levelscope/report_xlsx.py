"""요약 xlsx 생성: Levels / Entities(플러그인 제공 시) / Summary 시트."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BASE = Font(name="Arial", size=10)
NOTE = Font(name="Arial", size=9, italic=True, color="595959")
SUB_FILL = PatternFill("solid", fgColor="D6E4F0")
SUB = Font(name="Arial", size=10, bold=True)
_thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal="center")
CENTERM = Alignment(horizontal="center", vertical="center")


def _write_header(ws, cols, row=1):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row, j, c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTERM
    ws.row_dimensions[row].height = 28


def build_workbook(df, cfg, out_path, entities=None, notes=None, log=print):
    """df: Levels 시트용 DataFrame(set, level 선두). entities: (시트명, 헤더, 행들)"""
    cols = list(df.columns)
    N = len(df)
    wb = Workbook()
    ws = wb.active
    ws.title = "Levels"
    _write_header(ws, cols)
    text_cols = {c for c in cols if df[c].dtype == object} | {"set"}
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            cell = ws.cell(i, j, v)
            cell.font = BASE
            cell.border = BORDER
            if cols[j - 1] not in text_cols:
                cell.alignment = CENTER
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{N + 1}"
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = 30 if c == "set" else max(7, min(len(str(c)) + 2, 18))

    if entities:
        sheet_name, headers, rows = entities
        es = wb.create_sheet(sheet_name)
        _write_header(es, headers)
        r = 2
        for rec in rows:
            for j, v in enumerate(rec, 1):
                c = es.cell(r, j, v)
                c.font = BASE
            r += 1
        es.freeze_panes = "A2"
        es.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{r - 1}"
        for j, h in enumerate(headers, 1):
            es.column_dimensions[get_column_letter(j)].width = 30 if h == "set" else max(8, min(len(str(h)) + 4, 30))
        log(f"[xlsx] {sheet_name} 시트 {r - 2}행")

    # ---- Summary ----
    sm = wb.create_sheet("Summary")
    L = {c: get_column_letter(j) for j, c in enumerate(cols, 1)}
    D = lambda c: f"Levels!${L[c]}$2:${L[c]}${N + 1}"  # noqa: E731
    scfg = cfg.get("summary") or {}
    avg_cols = [c for c in (scfg.get("avg_cols") or []) if c in cols]
    r = 1
    sm.cell(r, 1, f"{cfg.get('game', '')} 레벨 데이터 요약 ({N} levels)").font = Font(
        name="Arial", size=12, bold=True, color="1F3864")
    r += 2

    sm.cell(r, 1, "세트별").font = SUB
    r += 1
    hdrs = ["set", "levels"] + [f"avg_{c}" for c in avg_cols]
    for j, h in enumerate(hdrs, 1):
        c = sm.cell(r, j, h); c.fill = SUB_FILL; c.font = SUB; c.border = BORDER
    r += 1
    sets = list(pd.unique(df["set"]))
    first = r
    for s in sets:
        sm.cell(r, 1, s)
        sm.cell(r, 2, f"=COUNTIF({D('set')},$A{r})")
        for k, c in enumerate(avg_cols):
            sm.cell(r, 3 + k, f"=ROUND(AVERAGEIFS({D(c)},{D('set')},$A{r}),2)")
        for j in range(1, len(hdrs) + 1):
            sm.cell(r, j).font = BASE; sm.cell(r, j).border = BORDER
        r += 1
    sm.cell(r, 1, "TOTAL").font = SUB
    sm.cell(r, 2, f"=SUM(B{first}:B{r - 1})")
    for j in range(1, len(hdrs) + 1):
        sm.cell(r, j).border = BORDER; sm.cell(r, j).font = SUB
    r += 2

    cat = scfg.get("category_col")
    if cat and cat in cols:
        order = scfg.get("category_order") or sorted(x for x in df[cat].dropna().unique())
        sm.cell(r, 1, f"{cat}별").font = SUB
        r += 1
        hdrs2 = [cat, "levels"] + [f"avg_{c}" for c in avg_cols]
        for j, h in enumerate(hdrs2, 1):
            c = sm.cell(r, j, h); c.fill = SUB_FILL; c.font = SUB; c.border = BORDER
        r += 1
        for dn in order:
            sm.cell(r, 1, dn)
            sm.cell(r, 2, f"=COUNTIF({D(cat)},$A{r})")
            for k, c in enumerate(avg_cols):
                sm.cell(r, 3 + k, f"=ROUND(AVERAGEIFS({D(c)},{D(cat)},$A{r}),2)")
            for j in range(1, len(hdrs2) + 1):
                sm.cell(r, j).font = BASE; sm.cell(r, j).border = BORDER
            r += 1
        r += 1

    count_cols = [c for c in (cfg.get("fields", {}).get("counts") or {}) if c in cols]
    if count_cols:
        sm.cell(r, 1, "요소(기믹)별 등장 현황").font = SUB
        r += 1
        for j, h in enumerate(["element", "levels_with (>0)", "total_count"], 1):
            c = sm.cell(r, j, h); c.fill = SUB_FILL; c.font = SUB; c.border = BORDER
        r += 1
        for g in count_cols:
            sm.cell(r, 1, g)
            sm.cell(r, 2, f'=COUNTIF({D(g)},">0")')
            sm.cell(r, 3, f"=SUM({D(g)})")
            for j in range(1, 4):
                sm.cell(r, j).font = BASE; sm.cell(r, j).border = BORDER
            r += 1
        r += 1
    for n in (notes or []):
        sm.cell(r, 1, n).font = NOTE
        r += 1
    for col, w in {"A": 36, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14}.items():
        sm.column_dimensions[col].width = w

    wb.save(out_path)
    log(f"[xlsx] 저장: {out_path} (Levels {N}행, 열 {len(cols)}개)")
